from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import (
    FormulaRule,
)
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)
from openpyxl.worksheet.datavalidation import (
    DataValidation,
)
from openpyxl.worksheet.table import (
    Table,
    TableStyleInfo,
)

from services.crm.models import CRMProject
from services.crm.storage import (
    EXPORT_ROOT,
    ensure_crm_directories,
    load_projects,
)


HEADERS = [
    "Project ID",
    "Client Name",
    "Company",
    "Ticker",
    "Industry",
    "Country",
    "Lead Source",
    "Package",
    "Project Type",
    "Status",
    "Priority",
    "Order Date",
    "Due Date",
    "Delivery Date",
    "Price (USD)",
    "Hours Worked",
    "Effective Hourly Rate",
    "Revision Count",
    "Client Rating",
    "Client Email",
    "Client Phone",
    "Website",
    "Primary Objective",
    "Intended Audience",
    "Requested Deliverables",
    "Requested Focus Areas",
    "Assigned Analyst",
    "Folder Path",
    "Payment Status",
    "Invoice #",
    "Follow-up Date",
    "Repeat Client",
    "Data Source Mode",
    "SEC Data Status",
    "Financial Data Status",
    "Notes",
]


def _row(
    project: CRMProject,
    excel_row: int,
) -> list:
    return [
        project.project_id,
        project.client_name,
        project.company_name,
        project.ticker,
        project.industry,
        project.country,
        project.lead_source,
        project.package,
        project.project_type,
        project.status,
        project.priority,
        project.order_date,
        project.due_date,
        project.delivery_date,
        project.price_usd,
        project.hours_worked,
        (
            f'=IF(P{excel_row}>0,'
            f'O{excel_row}/P{excel_row},"")'
        ),
        project.revision_count,
        project.client_rating,
        project.client_email,
        project.client_phone,
        project.company_website,
        project.primary_objective,
        project.intended_audience,
        project.requested_deliverables,
        project.requested_focus_areas,
        project.assigned_analyst,
        project.folder_path,
        project.payment_status,
        project.invoice_number,
        project.follow_up_date,
        "Yes" if project.repeat_client else "No",
        project.data_source_mode,
        project.sec_data_status,
        project.financial_data_status,
        project.notes,
    ]


def export_master_crm(
    output_path: str | Path | None = None,
) -> Path:
    ensure_crm_directories()

    projects = load_projects()

    if output_path is None:
        timestamp = datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )

        output_path = (
            EXPORT_ROOT
            / f"Veles_Master_CRM_{timestamp}.xlsx"
        )

    output = Path(output_path)
    output.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Client CRM"

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="0B162B",
    )

    for column, header in enumerate(
        HEADERS,
        start=1,
    ):
        cell = worksheet.cell(
            row=1,
            column=column,
            value=header,
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for excel_row, project in enumerate(
        projects,
        start=2,
    ):
        for column, value in enumerate(
            _row(project, excel_row),
            start=1,
        ):
            worksheet.cell(
                row=excel_row,
                column=column,
                value=value,
            )

    last_row = max(2, len(projects) + 1)
    last_column = worksheet.cell(
        row=1,
        column=len(HEADERS),
    ).column_letter

    table = Table(
        displayName="VelesClientCRM",
        ref=f"A1:{last_column}{last_row}",
    )

    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showRowStripes=True,
        showFirstColumn=False,
        showLastColumn=False,
    )

    worksheet.add_table(table)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        f"A1:{last_column}{last_row}"
    )

    widths = {
        "A": 16,
        "B": 22,
        "C": 28,
        "D": 12,
        "E": 20,
        "F": 16,
        "G": 15,
        "H": 14,
        "I": 22,
        "J": 16,
        "K": 12,
        "L": 14,
        "M": 14,
        "N": 14,
        "O": 14,
        "P": 14,
        "Q": 20,
        "R": 15,
        "S": 14,
        "T": 28,
        "U": 18,
        "V": 28,
        "W": 35,
        "X": 24,
        "Y": 35,
        "Z": 35,
        "AA": 20,
        "AB": 44,
        "AC": 16,
        "AD": 16,
        "AE": 16,
        "AF": 14,
        "AG": 24,
        "AH": 18,
        "AI": 22,
        "AJ": 45,
    }

    for column, width in widths.items():
        worksheet.column_dimensions[
            column
        ].width = width

    worksheet.column_dimensions[
        "O"
    ].number_format = "$#,##0.00"

    lists = workbook.create_sheet("Lists")

    options = {
        "A": [
            "Status",
            "Lead",
            "In Progress",
            "Waiting on Client",
            "Quality Control",
            "Delivered",
            "Completed",
            "Archived",
        ],
        "B": [
            "Lead Source",
            "Fiverr",
            "Website",
            "Referral",
            "LinkedIn",
            "Direct",
        ],
        "C": [
            "Package",
            "Basic",
            "Standard",
            "Premium",
            "Custom",
        ],
        "D": [
            "Priority",
            "Low",
            "Medium",
            "High",
            "Urgent",
        ],
        "E": [
            "Payment Status",
            "Pending",
            "Paid",
            "Partially Paid",
            "Refunded",
        ],
        "F": [
            "Project Type",
            "Equity Research",
            "Financial Modeling",
            "SEC Filing Review",
            "Competitive Analysis",
            "Biotechnology Intelligence",
            "Custom Research",
        ],
    }

    for column, values in options.items():
        for row_index, value in enumerate(
            values,
            start=1,
        ):
            lists[
                f"{column}{row_index}"
            ] = value

            if row_index == 1:
                lists[
                    f"{column}{row_index}"
                ].font = Font(bold=True)

    validation_rules = [
        ("J2:J1000", "'Lists'!$A$2:$A$8"),
        ("G2:G1000", "'Lists'!$B$2:$B$6"),
        ("H2:H1000", "'Lists'!$C$2:$C$5"),
        ("K2:K1000", "'Lists'!$D$2:$D$5"),
        ("AC2:AC1000", "'Lists'!$E$2:$E$5"),
        ("I2:I1000", "'Lists'!$F$2:$F$7"),
    ]

    for target_range, formula in validation_rules:
        validation = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=True,
        )

        worksheet.add_data_validation(
            validation
        )
        validation.add(target_range)

    red_fill = PatternFill(
        fill_type="solid",
        fgColor="F4CCCC",
    )

    worksheet.conditional_formatting.add(
        f"A2:{last_column}{last_row}",
        FormulaRule(
            formula=[
                'AND($M2<TODAY(),'
                '$J2<>"Completed",'
                '$M2<>"")'
            ],
            fill=red_fill,
        ),
    )

    workbook.save(output)

    return output
